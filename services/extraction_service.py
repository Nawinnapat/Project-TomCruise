from __future__ import annotations

from pathlib import Path

from models.company import CompanyInput, CompanyRecord
from parser.dbd_parser import parse_company_record, parse_search_results
from scraper.dbd_client import DbdDataWarehouseClient
from services.matching import best_match
from utils.logging import get_logger
from utils.validation import valid_juristic_id

processing_logger = get_logger("processing")
error_logger = get_logger("errors")


class ExtractionService:
    def __init__(self, client: DbdDataWarehouseClient | None = None) -> None:
        self.client = client or DbdDataWarehouseClient()

    def extract_many(self, inputs: list[CompanyInput], fiscal_years: list[str] | None = None) -> list[CompanyRecord]:
        records: list[CompanyRecord] = []
        for company_input in inputs:
            try:
                records.append(self.extract_one(company_input, fiscal_years=fiscal_years))
            except Exception as exc:  # noqa: BLE001 - batch must continue
                error_logger.error("company extraction failed input=%s error=%s", company_input, exc)
                records.append(
                    CompanyRecord(
                        company_name=company_input.company_name,
                        juristic_id=company_input.juristic_id,
                        extraction_timestamp=None,
                        errors=[str(exc)],
                    )
                )
        return records

    def extract_one(self, company_input: CompanyInput, fiscal_years: list[str] | None = None) -> CompanyRecord:
        juristic_id = company_input.juristic_id
        juristic_type = "5"
        confidence = 100 if valid_juristic_id(juristic_id) else 0
        search_result: dict[str, object] | None = None

        if juristic_id:
            search_payload = self.client.search_companies(juristic_id)
            for candidate in parse_search_results(search_payload):
                if candidate.get("juristic_id") == juristic_id:
                    juristic_type = str(candidate.get("juristic_type") or juristic_type)
                    raw_candidate = candidate.get("raw")
                    search_result = raw_candidate if isinstance(raw_candidate, dict) else None
                    break
        else:
            search_payload = self.client.search_companies(company_input.company_name or "")
            candidates = parse_search_results(search_payload)
            match, confidence = best_match(company_input, candidates)
            if not match:
                return CompanyRecord(company_name=company_input.company_name, confidence_score=0, errors=["NOT_FOUND"])
            juristic_id = str(match.get("juristic_id") or "")
            juristic_type = str(match.get("juristic_type") or "5")
            raw_candidate = match.get("raw")
            search_result = raw_candidate if isinstance(raw_candidate, dict) else None

        if not valid_juristic_id(juristic_id):
            return CompanyRecord(
                company_name=company_input.company_name,
                juristic_id=juristic_id,
                confidence_score=confidence,
                errors=["INVALID_JURISTIC_ID"],
            )

        profile = self.client.company_profile(juristic_type, juristic_id)
        history = self._history(juristic_type, juristic_id)
        years = fiscal_years or _profile_fiscal_years(profile)
        financials = self._financials(juristic_type, juristic_id, years)
        record = parse_company_record(profile, financials=financials, history=history, search_result=search_result)
        record.juristic_id = record.juristic_id or juristic_id
        record.confidence_score = confidence
        record.review_required = confidence < 60
        record.source_url = f"https://datawarehouse.dbd.go.th/company/profile/{juristic_type}{juristic_id}"
        processing_logger.info("company extracted juristic_id=%s confidence=%s", juristic_id, confidence)
        return record

    def _history(self, juristic_type: str, juristic_id: str) -> list[dict[str, object]]:
        history: list[dict[str, object]] = []
        for fetcher in (self.client.name_changes, self.client.capital_changes, self.client.mergers):
            try:
                payload = fetcher(juristic_type, juristic_id)
                history.append({"source": fetcher.__name__, "data": payload})
            except Exception as exc:  # noqa: BLE001
                error_logger.error("history endpoint failed juristic_id=%s error=%s", juristic_id, exc)
        return history

    def _financials(self, juristic_type: str, juristic_id: str, fiscal_years: list[str]) -> dict[str, object]:
        financials: dict[str, object] = {}
        for year in fiscal_years[:5]:
            try:
                financials[year] = self.client.financial_statements(juristic_type, juristic_id, year)
            except Exception as exc:  # noqa: BLE001
                error_logger.error("financial endpoint failed juristic_id=%s year=%s error=%s", juristic_id, year, exc)
        return financials


def load_inputs(path: Path) -> list[CompanyInput]:
    if path.suffix.lower() == ".csv":
        return _load_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _load_xlsx(path)
    raise ValueError(f"Unsupported input file: {path}")


def _load_csv(path: Path) -> list[CompanyInput]:
    import csv

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return [
            CompanyInput(company_name=row.get("company_name") or row.get("name"), juristic_id=row.get("juristic_id"))
            for row in csv.DictReader(file)
        ]


def _load_xlsx(path: Path) -> list[CompanyInput]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value else "" for value in rows[0]]
    results = []
    for row in rows[1:]:
        data = dict(zip(headers, row, strict=False))
        results.append(
            CompanyInput(
                company_name=data.get("company_name") or data.get("name"),
                juristic_id=str(data.get("juristic_id") or "") or None,
            )
        )
    return results


def _profile_fiscal_years(profile: object) -> list[str]:
    if not isinstance(profile, dict):
        return []
    year = profile.get("fiscalYear") or profile.get("fiscalYearRpt")
    if year is None:
        return []
    return [str(year)]
